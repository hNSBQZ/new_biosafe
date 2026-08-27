"""Idempotent importer for the historical Excel seed."""

from __future__ import annotations

import ast
import hashlib
import json
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from biosafe.domain.history import HistoryCreate
from biosafe.storage.history_repository import HistoryRepository

REQUIRED_HEADERS = {"timestamp", "question", "answer", "original_answer", "references"}


@dataclass
class ImportStats:
    added: int = 0
    skipped: int = 0
    errors: int = 0


def import_excel(path: Path, repository: HistoryRepository) -> ImportStats:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    if not REQUIRED_HEADERS.issubset(headers):
        missing = sorted(REQUIRED_HEADERS.difference(headers))
        raise ValueError(f"Excel is missing required columns: {', '.join(missing)}")
    indexes = {name: headers.index(name) for name in REQUIRED_HEADERS}
    stats = ImportStats()
    for row in rows:
        try:
            mapped = {name: row[indexes[name]] for name in REQUIRED_HEADERS}
            question = _text(mapped["question"])
            if not question:
                stats.errors += 1
                continue
            created_at = _timestamp(mapped["timestamp"])
            fields = {
                "question": question,
                "system_answer": _text(mapped["original_answer"]),
                "corrected_answer": _text(mapped["answer"]),
                "references": _references(mapped["references"]),
                "created_at": created_at,
            }
            fingerprint = hashlib.sha256(
                json.dumps(fields, ensure_ascii=False, sort_keys=True).encode("utf-8")
            ).hexdigest()
            if repository.exists_seed(fingerprint):
                stats.skipped += 1
                continue
            repository.create(
                HistoryCreate(
                    request_id=f"seed-{uuid.uuid4()}",
                    session_id="seed",
                    created_at=created_at,
                    experiment_id="",
                    input_mode="seed",
                    question=question,
                    answer_source="seed",
                    system_answer=fields["system_answer"],
                    corrected_answer=fields["corrected_answer"],
                    references=fields["references"],
                    status="completed",
                    seed_fingerprint=fingerprint,
                )
            )
            stats.added += 1
        except (ValueError, TypeError, sqlite3.DatabaseError):
            stats.errors += 1
    return stats


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _timestamp(value: Any) -> str:
    if isinstance(value, datetime):
        aware = value if value.tzinfo else value.replace(tzinfo=UTC)
        return aware.isoformat()
    text = _text(value)
    return text or datetime.now(UTC).isoformat()


def _references(value: Any) -> list[dict[str, Any]]:
    if value is None or not _text(value):
        return []
    if isinstance(value, (list, dict)):
        parsed = value
    else:
        raw = _text(value)
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            try:
                parsed = ast.literal_eval(raw)
            except (ValueError, SyntaxError):
                return [{"content": raw, "raw_metadata": {"seed_unparsed": True}}]
    if isinstance(parsed, dict):
        parsed = [parsed]
    if not isinstance(parsed, list):
        return [{"content": _text(parsed), "raw_metadata": {"seed_unparsed": True}}]
    return [item if isinstance(item, dict) else {"content": _text(item)} for item in parsed]
